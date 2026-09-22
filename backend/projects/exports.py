import csv

from django.http import HttpResponse

EXPORT_COLUMNS = [
    ("Project Id", lambda p: p.code),
    ("Project Name", lambda p: p.name),
    ("Percentage Completed", lambda p: p.percent_complete),
    ("Status", lambda p: p.get_status_display()),
    ("Owner", lambda p: p.owner.username if p.owner else ""),
    ("PMO Owner", lambda p: p.pmo.username if p.pmo else ""),
    ("Project Group", lambda p: p.project_group),
    ("Created By", lambda p: p.created_by.username if p.created_by else ""),
    ("Start Date", lambda p: p.start_date),
    ("End Date", lambda p: p.end_date),
    ("Completion Time", lambda p: p.completion_date),
    ("Unique Order Id", lambda p: p.unique_order_id),
    ("PO Status", lambda p: p.po_status),
    ("PO No.", lambda p: p.po_number),
    ("Currency", lambda p: p.currency),
    ("PO Value", lambda p: p.po_value),
    ("Project Man Days", lambda p: p.man_days),
    ("Sales Person Name", lambda p: p.sales_person.username if p.sales_person else ""),
    ("Sales Region", lambda p: p.region),
    ("Client Name", lambda p: p.client.name if p.client else ""),
    ("Project Category-A", lambda p: p.category_a),
    ("Project Category-B", lambda p: p.category_b),
    ("Contract Type", lambda p: p.contract_type),
    ("Comments", lambda p: p.comments),
    ("Project Comments", lambda p: p.project_comments),
    ("Overrun Comments", lambda p: p.overrun_comments),
    ("PO Date", lambda p: p.po_date),
    ("Billing Entity", lambda p: p.billing_entity),
    ("Budget Type", lambda p: p.budget_type),
    ("Working Emp", lambda p: p.working_emp),
    ("Created Time", lambda p: p.created_at.strftime("%Y-%m-%d %H:%M") if p.created_at else ""),
    ("Lob Head", lambda p: p.lob_head.username if p.lob_head else ""),
    ("Delivery Leads", lambda p: p.delivery_leads),
    ("Project Type", lambda p: p.get_project_type_display() if p.project_type else ""),
    ("Project Completion Status", lambda p: p.get_completion_status_display() if p.completion_status else ""),
]


def export_projects_csv(projects):
    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = 'attachment; filename="projects.csv"'
    writer = csv.writer(response)
    writer.writerow([label for label, _ in EXPORT_COLUMNS])
    for project in projects:
        writer.writerow([getter(project) if getter(project) is not None else "" for _, getter in EXPORT_COLUMNS])
    return response


def export_projects_xlsx(projects):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Projects"

    for col_index, (label, _) in enumerate(EXPORT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_index, value=label)
        cell.font = Font(bold=True)

    for row_index, project in enumerate(projects, start=2):
        for col_index, (_, getter) in enumerate(EXPORT_COLUMNS, start=1):
            value = getter(project)
            ws.cell(row=row_index, column=col_index, value=value if value is not None else "")

    for col_index in range(1, len(EXPORT_COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(col_index)].width = 20

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="projects.xlsx"'
    wb.save(response)
    return response
